import asyncio
import io
import json
import os
import re
import sys
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlparse

import requests
from flask import Flask, jsonify, render_template, request, send_file
from jinja2 import ChoiceLoader, FileSystemLoader
from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent
OLLAMA_HOST = os.environ.get("OLLAMA_HOST", "http://localhost:11434")

# index.html is expected in ./templates, but also works next to app.py
app = Flask(__name__, template_folder=str(BASE / "templates"))
app.jinja_loader = ChoiceLoader([FileSystemLoader(str(BASE / "templates")), FileSystemLoader(str(BASE))])


# ---------- MCP client ----------
async def _call_tool(name: str, arguments: dict):
    params = StdioServerParameters(
        command=sys.executable,
        args=[str(BASE / "mcp_server.py")],
        env={**os.environ},
    )
    async with stdio_client(params) as (read, write):
        async with ClientSession(read, write) as session:
            await session.initialize()
            result = await session.call_tool(name, arguments)
            text = result.content[0].text if result.content else ""
            if getattr(result, "is_error", getattr(result, "isError", False)):
                raise RuntimeError(text or "The agent returned an error.")
            data = json.loads(text)
            if isinstance(data, dict) and data.get("error"):
                raise RuntimeError(data["error"])
            return data


def _readable(e: Exception) -> str:
    while getattr(e, "exceptions", None):  # unwrap ExceptionGroup
        e = e.exceptions[0]
    return str(e)


# ---------- Page snapshot (real selectors for the Playwright script) ----------
SNAPSHOT_JS = r"""() => {
  const vis = e => { const r = e.getBoundingClientRect(), s = getComputedStyle(e);
    return r.width > 0 && r.height > 0 && s.visibility !== 'hidden' && s.display !== 'none'; };
  const clean = t => (t || '').replace(/\s+/g, ' ').trim().slice(0, 80);
  const labelFor = e => {
    if (e.id) { const l = document.querySelector('label[for="' + CSS.escape(e.id) + '"]'); if (l) return clean(l.innerText); }
    const p = e.closest('label'); return p ? clean(p.innerText) : '';
  };
  const sel = 'input, textarea, select, button, a[href], [role=button], [role=link], [role=tab], [role=checkbox]';
  const elements = [...document.querySelectorAll(sel)].filter(vis).slice(0, 60).map(e => {
    const clickable = ['BUTTON', 'A'].includes(e.tagName) || e.getAttribute('role');
    return {
      tag: e.tagName.toLowerCase(), type: e.getAttribute('type') || undefined,
      id: e.id || undefined, name: e.getAttribute('name') || undefined,
      testid: e.getAttribute('data-testid') || e.getAttribute('data-test') || undefined,
      role: e.getAttribute('role') || undefined,
      label: labelFor(e) || e.getAttribute('aria-label') || undefined,
      placeholder: e.getAttribute('placeholder') || undefined,
      text: clickable ? (clean(e.innerText) || undefined) : undefined,
      href: e.tagName === 'A' ? e.getAttribute('href') : undefined,
      required: e.required || undefined,
    };
  });
  return { title: document.title, url: location.href,
           headings: [...document.querySelectorAll('h1,h2,h3')].filter(vis).slice(0, 8).map(h => clean(h.innerText)),
           elements };
}"""


def _snapshot_with_browser(url: str) -> dict:
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch()
        try:
            page = browser.new_page()
            page.goto(url, wait_until="domcontentloaded", timeout=30000)
            try:
                page.wait_for_load_state("networkidle", timeout=8000)
            except Exception:
                pass  # chatty pages never go idle; what has loaded is enough
            return page.evaluate(SNAPSHOT_JS)
        finally:
            browser.close()


class _Collector(HTMLParser):
    """Fallback when no browser is available: read static HTML only."""

    def __init__(self):
        super().__init__()
        self.title, self.headings, self.elements, self.labels = "", [], [], {}
        self._in_title = False
        self._buf = None      # (kind, key, [text parts])
        self._open_el = None

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "title":
            self._in_title = True
        elif tag in ("h1", "h2", "h3"):
            self._buf = ("h", None, [])
        elif tag == "label":
            self._buf = ("label", a.get("for"), [])
        elif tag in ("input", "textarea", "select", "button", "a"):
            if tag == "input" and a.get("type") == "hidden":
                return
            if tag == "a" and not a.get("href"):
                return
            el = {"tag": tag, "type": a.get("type"), "id": a.get("id"), "name": a.get("name"),
                  "testid": a.get("data-testid") or a.get("data-test"), "role": a.get("role"),
                  "label": a.get("aria-label"), "placeholder": a.get("placeholder"),
                  "href": a.get("href") if tag == "a" else None,
                  "required": True if "required" in a else None}
            self.elements.append(el)
            if tag in ("button", "a"):
                self._open_el, self._text = el, []

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        if self._buf:
            self._buf[2].append(data)
        if self._open_el is not None:
            self._text.append(data)

    def handle_endtag(self, tag):
        if tag == "title":
            self._in_title = False
        elif tag in ("h1", "h2", "h3", "label") and self._buf:
            kind, key, parts = self._buf
            text = " ".join("".join(parts).split())[:80]
            if kind == "h" and text:
                self.headings.append(text)
            elif kind == "label" and key and text:
                self.labels[key] = text
            self._buf = None
        elif tag in ("button", "a") and self._open_el is not None:
            self._open_el["text"] = " ".join("".join(self._text).split())[:80] or None
            self._open_el = None


def _snapshot_static(url: str) -> dict:
    r = requests.get(url, timeout=20, headers={"User-Agent": "Mozilla/5.0"})
    r.raise_for_status()
    c = _Collector()
    c.feed(r.text)
    for el in c.elements:
        if not el.get("label") and el.get("id") in c.labels:
            el["label"] = c.labels[el["id"]]
    return {"title": c.title.strip(), "url": url, "headings": c.headings[:8],
            "elements": c.elements[:60]}


def _compact(snapshot: dict) -> str:
    snapshot["elements"] = [{k: v for k, v in el.items() if v not in (None, "", False)}
                            for el in snapshot["elements"]]
    text = json.dumps(snapshot, separators=(",", ":"), ensure_ascii=False)
    return text[:8000]


def _normalize_url(url: str) -> str:
    url = url.strip()
    if url and "://" not in url:
        local = re.match(r"^(localhost|127\.|0\.0\.0\.0|\[::1\])", url)
        url = ("http://" if local else "https://") + url
    p = urlparse(url)
    if p.scheme not in ("http", "https") or not p.netloc:
        raise ValueError("Please enter a valid web address, e.g. https://myapp.com or http://localhost:3000")
    return url


# ---------- Routes ----------
@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/models")
def models():
    try:
        r = requests.get(f"{OLLAMA_HOST}/api/tags", timeout=5)
        r.raise_for_status()
        return jsonify([m["name"] for m in r.json().get("models", [])])
    except Exception:
        return jsonify([])


@app.post("/api/generate")
def generate():
    body = request.get_json(silent=True) or {}
    requirements = (body.get("requirements") or "").strip()
    if not requirements:
        return jsonify(error="Please enter some requirements."), 400
    try:
        count = int(body.get("count") or 15)
        cases = asyncio.run(_call_tool("generate_test_cases", {
            "requirements": requirements, "model": body.get("model", ""), "count": count}))
        return jsonify(test_cases=cases)
    except Exception as e:  # surface a readable message to the UI
        return jsonify(error=_readable(e)), 500


@app.post("/api/playwright")
def playwright_script():
    body = request.get_json(silent=True) or {}
    cases = body.get("test_cases") or []
    if not cases:
        return jsonify(error="Select at least one test case."), 400
    try:
        url = _normalize_url(body.get("url") or "")
    except ValueError as e:
        return jsonify(error=str(e)), 400

    note = None
    try:
        snapshot = _snapshot_with_browser(url)
    except Exception as browser_err:
        try:
            snapshot = _snapshot_static(url)
            note = ("Read the page without a browser, so JavaScript-rendered elements may be missing. "
                    "Run `playwright install chromium` for best results.")
        except Exception:
            msg = str(browser_err).splitlines()[0] if str(browser_err) else "unknown error"
            return jsonify(error=f"Could not open {url}: {msg}"), 502

    if not snapshot.get("elements"):
        note = ((note + " ") if note else "") + \
            "No interactive elements were found on this page (is a login required?)."

    try:
        result = asyncio.run(_call_tool("generate_playwright_script", {
            "test_cases_json": json.dumps(cases),
            "page_snapshot": _compact(snapshot),
            "url": url,
            "model": body.get("model", ""),
        }))
    except Exception as e:
        return jsonify(error=_readable(e)), 500

    warnings = [w for w in (note, result.get("warning")) if w]
    return jsonify(script=result["script"], warning=" ".join(warnings) or None,
                   page_title=snapshot.get("title"), elements=len(snapshot.get("elements", [])))


@app.post("/api/download")
def download():
    cases = (request.get_json(silent=True) or {}).get("test_cases", [])
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Cases"
    headers = ["ID", "Title", "Type", "Priority", "Preconditions", "Steps", "Expected Result"]
    ws.append(headers)
    for c in cases:
        steps = "\n".join(f"{i}. {s}" for i, s in enumerate(c.get("steps", []), 1))
        ws.append([c.get("id"), c.get("title"), c.get("type"), c.get("priority"),
                   c.get("preconditions"), steps, c.get("expected_result")])

    fill = PatternFill("solid", fgColor="4F46E5")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    widths = [10, 36, 14, 10, 30, 50, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name="test_cases.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


if __name__ == "__main__":
    app.run(debug=True, port=5000)
